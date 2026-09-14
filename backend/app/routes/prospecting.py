from __future__ import annotations

import csv
import io
import json
import secrets
import unicodedata
from datetime import date, datetime, timezone
from typing import Any, Literal

from fastapi import APIRouter, Depends, File, HTTPException, Query, Request, UploadFile
from pydantic import BaseModel, Field
from sqlalchemy import func
from sqlalchemy.orm import Session

from app.core.config import settings
from app.database.session import get_db
from app.models.commercial import Lead
from app.models.prospecting import Prospect, ProspectingActivity, ProspectingCampaign, ProspectingSuppression
from app.models.user import User
from app.services.prospecting_service import (
    VALID_PROSPECT_STATUSES,
    discover_companies,
    distribute_daily_queue,
    generate_draft,
    get_or_create_campaign,
    json_list,
    normalize_casa_company,
    process_due_followups,
    reset_monthly_spend,
    send_prospect_email,
    sync_inbox,
)


router = APIRouter(prefix="/prospecting", tags=["prospecting"])


class CampaignUpdate(BaseModel):
    seller_ids: list[int] | None = None
    daily_per_seller: int | None = Field(default=None, ge=1, le=20)
    monthly_budget_cents: int | None = Field(default=None, ge=0, le=5000)
    status: Literal["active", "paused"] | None = None


class ProspectUpdate(BaseModel):
    email: str | None = None
    phone: str | None = None
    contact_name: str | None = None
    contact_role: str | None = None
    website: str | None = None
    assigned_to_id: int | None = None
    status: str | None = None


class ContactPayload(BaseModel):
    outcome: Literal["call_made", "interested", "no_interest", "note"]
    note: str = Field(min_length=1, max_length=4000)


class EmailPayload(BaseModel):
    subject: str = Field(min_length=1, max_length=255)
    body: str = Field(min_length=1, max_length=12000)


class SuppressionPayload(BaseModel):
    reason: str = Field(default="Solicitou não receber contatos", min_length=1, max_length=255)


def _prospect_dict(item: Prospect) -> dict[str, Any]:
    return {
        column.name: getattr(item, column.name)
        for column in item.__table__.columns
    } | {
        "score_reasons": json_list(item.score_reasons),
        "employee_evidence": json_list(item.employee_evidence),
    }


def _activity_dict(item: ProspectingActivity) -> dict[str, Any]:
    return {column.name: getattr(item, column.name) for column in item.__table__.columns}


@router.get("/config")
def config_status(db: Session = Depends(get_db)):
    campaign = get_or_create_campaign(db)
    reset_monthly_spend(campaign)
    db.commit()
    return {
        "source": {
            "provider": "Casa dos Dados",
            "configured": bool(settings.casa_dos_dados_api_key),
            "price_per_company_cents": 1,
        },
        "email": {
            "provider": "Hostinger",
            "address": settings.hostinger_email_address,
            "configured": bool(settings.hostinger_email_address and settings.hostinger_email_password),
            "dry_run": settings.prospecting_dry_run,
            "smtp": settings.hostinger_smtp_host,
            "imap": settings.hostinger_imap_host,
        },
        "scheduler": bool(settings.prospecting_cron_secret),
        "privacy_url": settings.prospecting_privacy_url,
        "campaign": _campaign_dict(campaign),
    }


def _campaign_dict(campaign: ProspectingCampaign) -> dict[str, Any]:
    return {
        "id": campaign.id,
        "name": campaign.name,
        "status": campaign.status,
        "seller_ids": [int(value) for value in json_list(campaign.seller_ids) if str(value).isdigit()],
        "daily_per_seller": campaign.daily_per_seller,
        "opened_min_months": campaign.opened_min_months,
        "opened_max_months": campaign.opened_max_months,
        "priority_min_months": campaign.priority_min_months,
        "priority_max_months": campaign.priority_max_months,
        "monthly_budget_cents": campaign.monthly_budget_cents,
        "spent_cents": campaign.spent_cents,
        "spent_month": campaign.spent_month,
        "last_discovery_at": campaign.last_discovery_at,
    }


@router.patch("/campaign")
def update_campaign(payload: CampaignUpdate, db: Session = Depends(get_db)):
    campaign = get_or_create_campaign(db)
    data = payload.model_dump(exclude_unset=True)
    seller_ids = data.pop("seller_ids", None)
    if seller_ids is not None:
        valid = {row.id for row in db.query(User).filter(User.id.in_(seller_ids), User.is_active.is_(True)).all()} if seller_ids else set()
        if len(valid) != len(set(seller_ids)):
            raise HTTPException(422, "Há usuários comerciais inválidos ou inativos.")
        if len(valid) > 3:
            raise HTTPException(422, "Selecione no máximo três usuários comerciais.")
        campaign.seller_ids = json.dumps(seller_ids)
    for key, value in data.items():
        setattr(campaign, key, value)
    db.commit()
    db.refresh(campaign)
    return _campaign_dict(campaign)


@router.get("/summary")
def summary(db: Session = Depends(get_db)):
    campaign = get_or_create_campaign(db)
    reset_monthly_spend(campaign)
    grouped = dict(db.query(Prospect.status, func.count(Prospect.id)).group_by(Prospect.status).all())
    month_start = date.today().replace(day=1)
    period_start = datetime.combine(month_start, datetime.min.time(), tzinfo=timezone.utc)

    def distinct_prospects(*kinds: str) -> int:
        return int(db.query(func.count(func.distinct(ProspectingActivity.prospect_id))).filter(
            ProspectingActivity.occurred_at >= period_start,
            ProspectingActivity.kind.in_(kinds),
        ).scalar() or 0)

    db.commit()
    return {
        "total": sum(grouped.values()),
        "statuses": grouped,
        "monthly": {
            "qualified": distinct_prospects("discovered", "imported"),
            "returns": distinct_prospects("email_received", "interested"),
            "opportunities": grouped.get("interested", 0) + grouped.get("emailing", 0),
            "demonstrations": distinct_prospects("demonstration"),
            "converted": distinct_prospects("converted"),
        },
        "goals": {"qualified": 300, "returns": 20, "opportunities": 8, "demonstrations": 5, "converted": 1},
        "budget": {"spent_cents": campaign.spent_cents, "limit_cents": campaign.monthly_budget_cents},
    }


@router.get("/prospects")
def list_prospects(
    status: str | None = Query(default=None),
    assigned_to_id: int | None = Query(default=None),
    search: str | None = Query(default=None, max_length=100),
    limit: int = Query(default=100, ge=1, le=300),
    db: Session = Depends(get_db),
):
    query = db.query(Prospect)
    if status:
        values = [value for value in status.split(",") if value in VALID_PROSPECT_STATUSES]
        if values:
            query = query.filter(Prospect.status.in_(values))
    if assigned_to_id:
        query = query.filter(Prospect.assigned_to_id == assigned_to_id)
    if search:
        pattern = f"%{search.strip()}%"
        query = query.filter(
            (Prospect.company_name.ilike(pattern))
            | (Prospect.trade_name.ilike(pattern))
            | (Prospect.cnpj.ilike(pattern))
        )
    items = query.order_by(Prospect.score.desc(), Prospect.created_at.desc()).limit(limit).all()
    return [_prospect_dict(item) for item in items]


@router.get("/activities")
def list_activities(
    prospect_id: str | None = None,
    direction: Literal["inbound", "outbound", "internal"] | None = None,
    limit: int = Query(default=100, ge=1, le=300),
    db: Session = Depends(get_db),
):
    query = db.query(ProspectingActivity)
    if prospect_id:
        query = query.filter(ProspectingActivity.prospect_id == prospect_id)
    if direction:
        query = query.filter(ProspectingActivity.direction == direction)
    return [_activity_dict(item) for item in query.order_by(ProspectingActivity.occurred_at.desc()).limit(limit).all()]


@router.post("/discover")
async def discover(limit: int = Query(default=45, ge=1, le=100), db: Session = Depends(get_db)):
    result = await discover_companies(db, requested=limit)
    result["distribution"] = distribute_daily_queue(db)
    return result


@router.post("/distribute")
def distribute(db: Session = Depends(get_db)):
    return distribute_daily_queue(db)


@router.patch("/prospects/{prospect_id}")
def update_prospect(prospect_id: str, payload: ProspectUpdate, db: Session = Depends(get_db)):
    prospect = db.get(Prospect, prospect_id)
    if not prospect:
        raise HTTPException(404, "Prospecto não encontrado.")
    data = payload.model_dump(exclude_unset=True)
    if data.get("status") and data["status"] not in VALID_PROSPECT_STATUSES:
        raise HTTPException(422, "Status de prospecto inválido.")
    if "assigned_to_id" in data and data["assigned_to_id"]:
        user = db.get(User, data["assigned_to_id"])
        if not user or not user.is_active:
            raise HTTPException(422, "Usuário responsável inválido.")
        data["assigned_to_name"] = user.full_name
    for key, value in data.items():
        setattr(prospect, key, value)
    db.commit()
    db.refresh(prospect)
    return _prospect_dict(prospect)


@router.post("/prospects/{prospect_id}/approve")
def approve_prospect(prospect_id: str, db: Session = Depends(get_db)):
    prospect = db.get(Prospect, prospect_id)
    if not prospect:
        raise HTTPException(404, "Prospecto não encontrado.")
    if prospect.status in {"suppressed", "rejected", "converted"}:
        raise HTTPException(409, "Este prospecto não pode ser aprovado no estado atual.")
    prospect.status = "approved"
    db.add(ProspectingActivity(prospect_id=prospect.id, kind="approved", body="Dados revisados e abordagem aprovada."))
    db.commit()
    return _prospect_dict(prospect)


@router.post("/prospects/{prospect_id}/contact")
def register_contact(prospect_id: str, payload: ContactPayload, db: Session = Depends(get_db)):
    prospect = db.get(Prospect, prospect_id)
    if not prospect:
        raise HTTPException(404, "Prospecto não encontrado.")
    now = datetime.now(timezone.utc)
    kind = payload.outcome
    if kind == "interested":
        prospect.status = "interested"
        prospect.contact_permission = True
        prospect.contact_permission_at = now
        prospect.contact_permission_source = "Interesse confirmado em contato comercial"
        prospect.next_action_at = None
    elif kind == "no_interest":
        prospect.status = "suppressed"
        prospect.contact_permission = False
        prospect.next_action_at = None
        db.add(ProspectingSuppression(cnpj=prospect.cnpj, email=prospect.email, phone=prospect.phone, reason=payload.note[:255]))
    elif kind == "call_made":
        prospect.status = "contacted"
    prospect.last_contact_at = now
    db.add(ProspectingActivity(
        prospect_id=prospect.id,
        kind=kind,
        channel="phone" if kind != "note" else "system",
        direction="outbound" if kind != "note" else "internal",
        body=payload.note,
        occurred_at=now,
    ))
    db.commit()
    return _prospect_dict(prospect)


@router.post("/prospects/{prospect_id}/draft")
async def draft_prospect(prospect_id: str, db: Session = Depends(get_db)):
    prospect = db.get(Prospect, prospect_id)
    if not prospect:
        raise HTTPException(404, "Prospecto não encontrado.")
    return await generate_draft(prospect)


@router.post("/prospects/{prospect_id}/send-email")
async def email_prospect(prospect_id: str, payload: EmailPayload, db: Session = Depends(get_db)):
    prospect = db.get(Prospect, prospect_id)
    if not prospect:
        raise HTTPException(404, "Prospecto não encontrado.")
    return await send_prospect_email(db, prospect, payload.subject, payload.body)


@router.post("/prospects/{prospect_id}/demonstration")
def register_demonstration(prospect_id: str, note: str = Query(default="Demonstração agendada", max_length=500), db: Session = Depends(get_db)):
    prospect = db.get(Prospect, prospect_id)
    if not prospect:
        raise HTTPException(404, "Prospecto não encontrado.")
    prospect.status = "interested"
    prospect.contact_permission = True
    prospect.contact_permission_at = datetime.now(timezone.utc)
    prospect.contact_permission_source = "Demonstração solicitada ou agendada"
    db.add(ProspectingActivity(prospect_id=prospect.id, kind="demonstration", channel="meeting", direction="outbound", body=note))
    db.commit()
    return _prospect_dict(prospect)


@router.post("/prospects/{prospect_id}/convert")
def convert_to_lead(prospect_id: str, db: Session = Depends(get_db)):
    prospect = db.get(Prospect, prospect_id)
    if not prospect:
        raise HTTPException(404, "Prospecto não encontrado.")
    if prospect.converted_lead_id:
        return {"lead_id": prospect.converted_lead_id, "already_converted": True}
    lead = Lead(
        name=prospect.contact_name or prospect.trade_name or prospect.company_name,
        company=prospect.trade_name or prospect.company_name,
        email=prospect.email,
        phone=prospect.phone,
        status="qualificado",
        stage="qualificado",
        origin="Prospecção",
        notes=f"CNPJ: {prospect.cnpj}. Origem: {prospect.source}. Score: {prospect.score}/100.",
        next_action="Agendar diagnóstico e demonstração",
        next_contact_date=date.today(),
        assigned_to_id=prospect.assigned_to_id,
        assigned_to_name=prospect.assigned_to_name,
    )
    db.add(lead)
    db.flush()
    prospect.converted_lead_id = lead.id
    prospect.status = "converted"
    db.add(ProspectingActivity(prospect_id=prospect.id, kind="converted", body="Convertido para o pipeline comercial."))
    db.commit()
    return {"lead_id": lead.id, "already_converted": False}


@router.post("/prospects/{prospect_id}/suppress")
def suppress_prospect(prospect_id: str, payload: SuppressionPayload, db: Session = Depends(get_db)):
    prospect = db.get(Prospect, prospect_id)
    if not prospect:
        raise HTTPException(404, "Prospecto não encontrado.")
    prospect.status = "suppressed"
    prospect.contact_permission = False
    prospect.next_action_at = None
    db.add(ProspectingSuppression(cnpj=prospect.cnpj, email=prospect.email, phone=prospect.phone, reason=payload.reason))
    db.add(ProspectingActivity(prospect_id=prospect.id, kind="suppressed", body=payload.reason))
    db.commit()
    return _prospect_dict(prospect)


@router.post("/inbox/sync")
async def sync_hostinger_inbox(db: Session = Depends(get_db)):
    return await sync_inbox(db)


def _normalized_header(value: Any) -> str:
    text = unicodedata.normalize("NFKD", str(value or "")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", "_", text).strip("_")


def _read_upload(content: bytes, filename: str) -> list[dict[str, Any]]:
    if filename.lower().endswith(".xlsx"):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise HTTPException(500, "Suporte a Excel não está instalado no backend.") from exc
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [_normalized_header(value) for value in rows[0]]
        return [dict(zip(headers, row)) for row in rows[1:] if any(value not in (None, "") for value in row)]
    try:
        decoded = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        decoded = content.decode("latin-1")
    if not decoded.strip():
        return []
    sample = decoded[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;	")
    except csv.Error:
        dialect = csv.excel
    return [{_normalized_header(key): value for key, value in row.items()} for row in csv.DictReader(io.StringIO(decoded), dialect=dialect)]


def _pick(row: dict[str, Any], *names: str) -> Any:
    return next((row[name] for name in names if row.get(name) not in (None, "")), None)


@router.post("/import")
async def import_prospects(file: UploadFile = File(...), db: Session = Depends(get_db)):
    filename = file.filename or ""
    if not filename.lower().endswith((".csv", ".xlsx")):
        raise HTTPException(422, "Envie um arquivo CSV ou XLSX.")
    content = await file.read()
    if len(content) > 10 * 1024 * 1024:
        raise HTTPException(413, "O arquivo deve ter no máximo 10 MB.")
    rows = _read_upload(content, filename)
    campaign = get_or_create_campaign(db)
    added = 0
    skipped = 0
    for row in rows[:5000]:
        raw = {
            "cnpj": _pick(row, "cnpj", "documento"),
            "razao_social": _pick(row, "razao_social", "empresa", "company"),
            "nome_fantasia": _pick(row, "nome_fantasia", "fantasia"),
            "codigo_atividade_principal": _pick(row, "cnae", "cnae_principal", "codigo_atividade_principal"),
            "descricao_atividade_principal": _pick(row, "atividade", "ramo", "cnae_descricao"),
            "data_abertura": _pick(row, "data_abertura", "abertura"),
            "porte_empresa": _pick(row, "porte", "porte_empresa"),
            "capital_social": _pick(row, "capital_social", "capital") or 0,
            "email": _pick(row, "email", "e_mail"),
            "telefone": _pick(row, "telefone", "phone", "celular"),
            "municipio": _pick(row, "municipio", "cidade"),
            "uf": _pick(row, "uf", "estado"),
        }
        data = normalize_casa_company(raw)
        if not data["cnpj"] or db.query(Prospect).filter(Prospect.cnpj == data["cnpj"]).first():
            skipped += 1
            continue
        clean = {key: value for key, value in data.items() if key != "is_mei"}
        clean["score_reasons"] = json.dumps(clean["score_reasons"], ensure_ascii=False)
        clean["employee_evidence"] = json.dumps(clean["employee_evidence"], ensure_ascii=False)
        prospect = Prospect(campaign_id=campaign.id, source=f"Importação: {filename[:60]}", **clean)
        db.add(prospect)
        db.flush()
        db.add(ProspectingActivity(prospect_id=prospect.id, kind="imported", body=f"Importado do arquivo {filename[:120]}."))
        added += 1
    db.commit()
    distribution = distribute_daily_queue(db)
    return {"rows": min(len(rows), 5000), "added": added, "skipped": skipped, "distribution": distribution}


@router.post("/run")
async def scheduled_run(request: Request, db: Session = Depends(get_db)):
    authorization = request.headers.get("authorization", "")
    expected = f"Bearer {settings.prospecting_cron_secret}"
    if not settings.prospecting_cron_secret or not secrets.compare_digest(authorization, expected):
        raise HTTPException(401, "Cron não autorizado.")
    campaign = get_or_create_campaign(db)
    if campaign.status != "active":
        return {"skipped": True, "reason": "Campanha pausada."}
    inbox = await sync_inbox(db) if settings.hostinger_email_password else {"added": 0, "unmatched": 0}
    followups = await process_due_followups(db)
    discovery: dict[str, Any] = {"skipped": True, "reason": "Fonte ainda não configurada."}
    if settings.casa_dos_dados_api_key:
        discovery = await discover_companies(db, requested=max(15, campaign.daily_per_seller * 9))
    distribution = distribute_daily_queue(db)
    return {"inbox": inbox, "followups": followups, "discovery": discovery, "distribution": distribution}


@router.post("/inbox/run")
async def scheduled_inbox_sync(request: Request, db: Session = Depends(get_db)):
    authorization = request.headers.get("authorization", "")
    expected = f"Bearer {settings.prospecting_cron_secret}"
    if not settings.prospecting_cron_secret or not secrets.compare_digest(authorization, expected):
        raise HTTPException(401, "Cron não autorizado.")
    inbox = await sync_inbox(db) if settings.hostinger_email_password else {"added": 0, "unmatched": 0}
    followups = await process_due_followups(db)
    return {"inbox": inbox, "followups": followups}
